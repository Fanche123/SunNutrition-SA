import argparse
import json
import os
import re
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula


ROOT_DIR = Path(__file__).resolve().parents[1]
REGISTRY_PATH = ROOT_DIR / "backend" / "table-registry.json"
DEFAULT_OUTPUT = ROOT_DIR / "tmp" / "backend-data-cache.json"
BACKEND_TABLE_PROJECTIONS = {
    "etiquetas": {
        "columns": ["id_etiqueta", "etiqueta", "categoria_pnl"],
        "aliases": {
            "id_etiqueta": ["id_etiqueta"],
            "etiqueta": ["etiqueta", "etiquetas", "nombre"],
            "categoria_pnl": ["categoria_pnl", "categoria_p_n_l", "categoria", "categoria_estado_resultados"],
        },
    },
    "acreedores": {
        "columns": ["id_acreedor", "origen_tipo_acreedor", "origen_id_acreedor", "cuit_cuil", "cbu_alias", "acuerdo_de_pago", "etiqueta"],
        "aliases": {
            "id_acreedor": ["id_acreedor"],
            "origen_tipo_acreedor": ["origen_tipo_acreedor", "tipo"],
            "origen_id_acreedor": ["origen_id_acreedor", "id_proveedor", "id_empleado", "id_canal", "id_flete"],
            "cuit_cuil": ["cuit_cuil", "cuit", "cuil"],
            "cbu_alias": ["cbu_alias", "cbu", "alias"],
            "acuerdo_de_pago": ["acuerdo_de_pago"],
            "etiqueta": ["etiqueta", "categoria_gasto", "tipo_gasto", "id_etiqueta"],
        },
    },
    "acreedores_etiquetas": {
        "columns": ["id_acreedor_etiqueta", "id_acreedor", "id_etiqueta"],
        "aliases": {
            "id_acreedor_etiqueta": ["id_acreedor_etiqueta", "id"],
            "id_acreedor": ["id_acreedor"],
            "id_etiqueta": ["id_etiqueta"],
        },
    },
    "proveedores": {
        "columns": ["id_proveedor", "nombre", "telefono", "email", "tiempo_estimado_entrega", "pedido_minimo"],
        "aliases": {
            "id_proveedor": ["id_proveedor"],
            "nombre": ["nombre"],
            "telefono": ["telefono"],
            "email": ["email"],
            "tiempo_estimado_entrega": ["tiempo_estimado_entrega"],
            "pedido_minimo": ["pedido_minimo", "minimo_de_pedido"],
        },
    },
    "clientes": {
        "columns": ["id_cliente", "nombre_cliente", "cuit", "direccion", "localidad", "contacto", "telefono", "email", "id_canal", "tipo", "tipo_comprobante", "plazo_cobro"],
        "aliases": {
            "id_cliente": ["id_cliente"],
            "nombre_cliente": ["nombre_cliente", "nombre"],
            "cuit": ["cuit"],
            "direccion": ["direccion"],
            "localidad": ["localidad"],
            "contacto": ["contacto", "nombre_contacto"],
            "telefono": ["telefono"],
            "email": ["email"],
            "id_canal": ["id_canal"],
            "tipo": ["tipo", "id_tipo_cliente"],
            "tipo_comprobante": ["tipo_comprobante", "tipo_de_comprobante"],
            "plazo_cobro": ["plazo_cobro", "id_plazo_cobro"],
        },
    },
    "empleados": {
        "columns": ["id_empleado", "nombre_empleado", "categoria_empleado", "contratacion", "fecha_alta", "fecha_baja", "dni", "fecha_nacimiento", "direccion_empleado", "localidad_empleado"],
        "aliases": {
            "id_empleado": ["id_empleado"],
            "nombre_empleado": ["nombre_empleado", "nombre"],
            "categoria_empleado": ["categoria_empleado", "categoria"],
            "contratacion": ["contratacion"],
            "fecha_alta": ["fecha_alta"],
            "fecha_baja": ["fecha_baja"],
            "dni": ["dni"],
            "fecha_nacimiento": ["fecha_nacimiento"],
            "direccion_empleado": ["direccion_empleado", "direccion"],
            "localidad_empleado": ["localidad_empleado", "localidad"],
        },
    },
    "canales": {
        "columns": ["id_canal", "nombre", "comision"],
        "aliases": {
            "id_canal": ["id_canal"],
            "nombre": ["nombre"],
            "comision": ["comision"],
        },
    },
    "fletes": {
        "columns": ["id_flete", "nombre_flete", "telefono_flete", "email_flete"],
        "aliases": {
            "id_flete": ["id_flete"],
            "nombre_flete": ["nombre_flete", "nombre"],
            "telefono_flete": ["telefono_flete", "telefono"],
            "email_flete": ["email_flete", "email"],
        },
    },
    "datos_bancarios": {
        "columns": ["id_dato_bancario", "id_acreedor", "tipo", "banco", "fecha", "concepto", "detalle", "nombre", "cuit", "cbu_alias", "tipo_doc", "nro_doc", "debito", "credito", "saldo", "canal"],
        "aliases": {
            "id_dato_bancario": ["id_dato_bancario", "id"],
            "id_acreedor": ["id_acreedor"],
            "tipo": ["tipo", "categoria", "categoria_gasto"],
            "banco": ["banco"],
            "fecha": ["fecha", "fecha_contable", "fecha_banco"],
            "concepto": ["concepto"],
            "detalle": ["detalle", "informacion_complementaria", "descripcion"],
            "nombre": ["nombre", "razon_social", "acreedor"],
            "cuit": ["cuit", "cuil"],
            "cbu_alias": ["cbu_alias", "cbu", "alias"],
            "tipo_doc": ["tipo_doc", "tipo_documento"],
            "nro_doc": ["nro_doc", "numero_doc", "documento"],
            "debito": ["debito", "debitos"],
            "credito": ["credito", "creditos"],
            "saldo": ["saldo"],
            "canal": ["canal"],
        },
    },
    "movimientos_bancarios": {
        "columns": ["id_movimiento_bancario", "banco", "fecha", "cod_concepto", "concepto", "detalle", "cuit", "nro_cheque", "debito", "credito", "importe", "saldo", "id_pago", "id_cobro"],
        "aliases": {
            "id_movimiento_bancario": ["id_movimiento_bancario", "id", "id_movimiento"],
            "banco": ["banco"],
            "fecha": ["fecha", "fecha_contable", "fecha_banco"],
            "cod_concepto": ["cod_concepto", "cod_de_concepto", "codigo_concepto"],
            "concepto": ["concepto"],
            "detalle": ["detalle", "informacion_complementaria", "descripcion"],
            "cuit": ["cuit", "cuil", "nro_doc", "nro_documento"],
            "nro_cheque": ["nro_cheque", "nro_de_cheque", "numero_cheque", "cheque"],
            "debito": ["debito", "debito_en", "debito_en_$", "debitos"],
            "credito": ["credito", "credito_en", "credito_en_$", "creditos"],
            "importe": ["importe", "monto", "movimiento"],
            "saldo": ["saldo", "saldo_en", "saldo_en_$"],
            "id_pago": ["id_pago"],
            "id_cobro": ["id_cobro"],
        },
    },
    "items": {
        "columns": ["id_item", "origen_tipo", "id_origen", "ud_conteo"],
        "aliases": {
            "id_item": ["id_item"],
            "origen_tipo": ["origen_tipo", "tipo"],
            "id_origen": ["id_origen"],
            "ud_conteo": ["ud_conteo"],
        },
    },
    "productos": {
        "columns": ["id_producto", "id_item", "nombre_producto", "capacidad_palet", "cantidad_individual", "precio_base", "costo_base", "costo_ud_individual"],
        "aliases": {
            "id_producto": ["id_producto"],
            "id_item": ["id_item"],
            "nombre_producto": ["nombre_producto", "nombre"],
            "capacidad_palet": ["capacidad_palet"],
            "cantidad_individual": ["cantidad_individual"],
            "precio_base": ["precio_base", "precio_original"],
            "costo_base": ["costo_base", "costo"],
            "costo_ud_individual": ["costo_ud_individual"],
        },
    },
    "subproductos": {
        "columns": ["id_subproducto", "id_item", "nombre_subproducto", "ud_conteo"],
        "aliases": {
            "id_subproducto": ["id_subproducto", "id_item"],
            "id_item": ["id_item"],
            "nombre_subproducto": ["nombre_subproducto", "nombre_producto", "nombre"],
            "ud_conteo": ["ud_conteo"],
        },
    },
    "insumos": {
        "columns": ["id_insumo", "nombre", "ud_receta", "cantidad_receta", "tipo"],
        "aliases": {
            "id_insumo": ["id_insumo"],
            "nombre": ["nombre"],
            "ud_receta": ["ud_receta"],
            "cantidad_receta": ["cantidad_receta"],
            "tipo": ["tipo"],
        },
    },
    "insumos_proveedores": {
        "columns": ["id_insumos_proveedores", "id_insumo", "id_proveedor", "ud_proveedor", "cantidad_proveedor", "precio", "iva"],
        "aliases": {
            "id_insumos_proveedores": ["id_insumos_proveedores"],
            "id_insumo": ["id_insumo"],
            "id_proveedor": ["id_proveedor"],
            "ud_proveedor": ["ud_proveedor"],
            "cantidad_proveedor": ["cantidad_proveedor", "cantidad_receta"],
            "precio": ["precio", "precio_unitario"],
            "iva": ["iva"],
        },
    },
    "recetas": {
        "columns": ["id_receta", "id_item_resultado", "id_item_componente", "cantidad_componente", "unidad_base"],
        "aliases": {
            "id_receta": ["id_receta", "_rowNumber"],
            "id_item_resultado": ["id_item_resultado"],
            "id_item_componente": ["id_item_componente", "id_item"],
            "cantidad_componente": ["cantidad_componente", "cantidad"],
            "unidad_base": ["unidad_base", "ud_receta"],
        },
    },
    "pedidos": {
        "columns": ["id_pedido", "fecha_pedido", "id_cliente", "fecha_entrega", "fecha_original"],
        "aliases": {
            "id_pedido": ["id_pedido"],
            "fecha_pedido": ["fecha_pedido", "fecha"],
            "id_cliente": ["id_cliente"],
            "fecha_entrega": ["fecha_entrega"],
            "fecha_original": ["fecha_original"],
        },
    },
    "detalle_pedidos": {
        "columns": ["id_detalle_pedido", "id_pedido", "id_producto", "cantidad_cajas", "impuesto", "precio_ud", "bonificacion"],
        "aliases": {
            "id_detalle_pedido": ["id_detalle_pedido"],
            "id_pedido": ["id_pedido"],
            "id_producto": ["id_producto"],
            "cantidad_cajas": ["cantidad_cajas"],
            "impuesto": ["impuesto"],
            "precio_ud": ["precio_ud", "precio_unidad"],
            "bonificacion": ["bonificacion", "bonif"],
        },
    },
    "entregas": {
        "columns": ["id_entrega", "fecha", "id_flete", "id_acreedor_etiqueta", "id_egreso"],
        "aliases": {
            "id_entrega": ["id_entrega"],
            "fecha": ["fecha"],
            "id_flete": ["id_flete"],
            "id_acreedor_etiqueta": ["id_acreedor_etiqueta"],
            "id_egreso": ["id_egreso"],
        },
    },
    "entregas_detalle": {
        "columns": ["id_entregas_detalle", "id_entrega", "id_pedido"],
        "aliases": {
            "id_entregas_detalle": ["id_entregas_detalle", "id_entrega_pedido", "_rowNumber"],
            "id_entrega": ["id_entrega"],
            "id_pedido": ["id_pedido"],
        },
    },
    "ventas": {
        "columns": ["id_venta", "id_pedido", "id_cliente", "id_entrega", "tipo_factura", "nro_factura", "fecha_factura", "fecha_acordada", "iva", "subtotal", "total"],
        "aliases": {
            "id_venta": ["id_venta"],
            "id_pedido": ["id_pedido"],
            "id_cliente": ["id_cliente"],
            "id_entrega": ["id_entrega"],
            "tipo_factura": ["tipo_factura"],
            "nro_factura": ["nro_factura"],
            "fecha_factura": ["fecha_factura"],
            "fecha_acordada": ["fecha_acordada", "fecha_prevista_cobro", "fecha_prevista_de_cobro"],
            "iva": ["iva"],
            "subtotal": ["subtotal", "sub_total"],
            "total": ["total"],
        },
    },
    "cobros": {
        "columns": ["id_cobro", "fecha_cobro", "metodo", "id_cliente", "monto", "banco"],
        "aliases": {
            "id_cobro": ["id_cobro"],
            "fecha_cobro": ["fecha_cobro", "fecha"],
            "metodo": ["metodo", "forma_cobro", "id_forma_pago"],
            "id_cliente": ["id_cliente"],
            "monto": ["monto", "total"],
            "banco": ["banco"],
        },
    },
    "cobros_detalle": {
        "columns": ["id_cobros_detalle", "id_cobro", "id_venta", "monto_cancelado"],
        "aliases": {
            "id_cobros_detalle": ["id_cobros_detalle", "id_cobro_venta", "_rowNumber"],
            "id_cobro": ["id_cobro"],
            "id_venta": ["id_venta"],
            "monto_cancelado": ["monto_cancelado"],
        },
    },
    "retenciones_ganancias": {
        "columns": ["id_retencion_ganancias", "id_cobro", "fecha", "id_cliente", "monto"],
        "aliases": {
            "id_retencion_ganancias": ["id_retencion_ganancias", "id_retencion", "_rowNumber"],
            "id_cobro": ["id_cobro"],
            "fecha": ["fecha", "fecha_retencion"],
            "id_cliente": ["id_cliente"],
            "monto": ["monto", "retencion_ganancias", "ret_ganancias"],
        },
    },
    "retenciones_iibb": {
        "columns": ["id_retencion_iibb", "id_cobro", "fecha", "id_cliente", "monto"],
        "aliases": {
            "id_retencion_iibb": ["id_retencion_iibb", "id_retencion", "_rowNumber"],
            "id_cobro": ["id_cobro"],
            "fecha": ["fecha", "fecha_retencion"],
            "id_cliente": ["id_cliente"],
            "monto": ["monto", "retencion_iibb", "ret_iibb", "ingresos_brutos"],
        },
    },
    "caja": {
        "columns": ["cuenta", "monto"],
        "aliases": {
            "cuenta": ["cuenta", "formato", "banco", "account"],
            "monto": ["monto", "saldo", "saldo_inicial"],
        },
    },
    "cheques_entregados": {
        "columns": ["id_cheque_entregado", "id_pago", "fecha_entregado", "monto", "acreedor", "id_acreedor", "nro_cheque", "fecha_uso", "estado"],
        "aliases": {
            "id_cheque_entregado": ["id_cheque_entregado", "id_chequeentregado"],
            "id_pago": ["id_pago"],
            "fecha_entregado": ["fecha_entregado"],
            "monto": ["monto"],
            "acreedor": ["acreedor"],
            "id_acreedor": ["id_acreedor"],
            "nro_cheque": ["nro_cheque"],
            "fecha_uso": ["fecha_uso", "fecha_de_uso"],
            "estado": ["estado"],
        },
    },
    "cheques_recibidos": {
        "columns": ["id_cheque_recibido", "id_cobro", "fecha_entregado", "monto", "cliente", "id_cliente", "nro_cheque", "fecha_uso", "estado", "banco"],
        "aliases": {
            "id_cheque_recibido": ["id_cheque_recibido", "id_chequerecibido"],
            "id_cobro": ["id_cobro"],
            "fecha_entregado": ["fecha_entregado"],
            "monto": ["monto"],
            "cliente": ["cliente"],
            "id_cliente": ["id_cliente"],
            "nro_cheque": ["nro_cheque"],
            "fecha_uso": ["fecha_uso", "fecha_de_uso"],
            "estado": ["estado"],
            "banco": ["banco"],
        },
    },
    "compras": {
        "columns": ["id_compra", "id_proveedor", "fecha_pedido", "fecha_entrega_prevista"],
        "aliases": {
            "id_compra": ["id_compra"],
            "id_proveedor": ["id_proveedor"],
            "fecha_pedido": ["fecha_pedido"],
            "fecha_entrega_prevista": ["fecha_entrega_prevista"],
        },
    },
    "detalle_compras": {
        "columns": [
            "id_detalle_compra",
            "id_compra",
            "id_insumos_proveedores",
            "cantidad",
        ],
        "aliases": {
            "id_detalle_compra": ["id_detalle_compra"],
            "id_compra": ["id_compra"],
            "id_insumos_proveedores": ["id_insumos_proveedores", "id_insumo", "id_item"],
            "cantidad": ["cantidad"],
        },
    },
    "recepciones": {
        "columns": ["id_recepcion", "fecha_recepcion", "id_empleado", "id_compra", "id_acreedor_etiqueta", "id_egreso", "archivo_recepcion"],
        "aliases": {
            "id_recepcion": ["id_recepcion"],
            "fecha_recepcion": ["fecha_recepcion", "fecha"],
            "id_empleado": ["id_empleado"],
            "id_compra": ["id_compra"],
            "id_acreedor_etiqueta": ["id_acreedor_etiqueta"],
            "id_egreso": ["id_egreso"],
            "archivo_recepcion": ["archivo_recepcion", "archivo", "comprobante"],
        },
    },
    "detalle_recepciones": {
        "columns": ["id_detalle_recepcion", "id_recepcion", "id_insumo", "cantidad_recibida"],
        "aliases": {
            "id_detalle_recepcion": ["id_detalle_recepcion"],
            "id_recepcion": ["id_recepcion"],
            "id_insumo": ["id_insumo", "id_item"],
            "cantidad_recibida": ["cantidad_recibida", "cantidad"],
        },
    },
    "otros_gastos": {
        "columns": ["id_otros_gastos", "fecha_otros_gastos", "id_acreedor", "id_acreedor_etiqueta", "id_egreso", "detalle"],
        "aliases": {
            "id_otros_gastos": ["id_otros_gastos", "id_gasto"],
            "fecha_otros_gastos": ["fecha_otros_gastos", "fecha"],
            "id_acreedor": ["id_acreedor"],
            "id_acreedor_etiqueta": ["id_acreedor_etiqueta"],
            "id_egreso": ["id_egreso"],
            "detalle": ["detalle"],
        },
    },
    "egresos": {
        "columns": ["id_egreso", "fecha_factura", "fecha_prevista_pago", "id_etiqueta", "tipo_factura", "nro_factura", "iva", "per_ret_iva", "per_ret_iibb", "imp_internos", "subtotal", "total"],
        "aliases": {
            "id_egreso": ["id_egreso"],
            "fecha_factura": ["fecha_factura"],
            "fecha_prevista_pago": ["fecha_prevista_pago", "fecha_prevista_de_pago", "fecha prevista pago", "fecha prevista de pago"],
            "id_etiqueta": ["id_etiqueta", "tipo", "categoria", "categoria_gasto"],
            "tipo_factura": ["tipo_factura", "tipo_de_factura"],
            "nro_factura": ["nro_factura"],
            "iva": ["iva"],
            "per_ret_iva": ["per_ret_iva"],
            "per_ret_iibb": ["per_ret_iibb"],
            "imp_internos": ["imp_internos", "i_internos"],
            "subtotal": ["subtotal"],
            "total": ["total"],
        },
    },
    "pagos": {
        "columns": ["id_pago", "fecha_pago", "metodo", "banco", "monto"],
        "aliases": {
            "id_pago": ["id_pago"],
            "fecha_pago": ["fecha_pago", "fecha"],
            "metodo": ["metodo", "id_forma_pago", "forma_pago"],
            "banco": ["banco"],
            "monto": ["monto"],
        },
    },
    "detalle_pagos": {
        "columns": ["id_detalle_pago", "id_pago", "id_egreso", "monto_cancelado"],
        "aliases": {
            "id_detalle_pago": ["id_detalle_pago", "id_pago_egreso"],
            "id_pago": ["id_pago"],
            "id_egreso": ["id_egreso"],
            "monto_cancelado": ["monto_cancelado"],
        },
    },
    "sueldos": {
        "columns": ["id_sueldo", "fecha", "id_empleado", "id_acreedor_etiqueta", "id_egreso", "valor_remunerativo", "valor_no_remunerativo", "hs_trabajadas", "hs_con_justificacion_medica", "hs_feriado", "hs_extra", "premios", "sueldo_bruto", "sueldo_neto"],
        "aliases": {
            "id_sueldo": ["id_sueldo", "_rowNumber"],
            "fecha": ["fecha", "mes"],
            "id_empleado": ["id_empleado"],
            "id_acreedor_etiqueta": ["id_acreedor_etiqueta"],
            "id_egreso": ["id_egreso"],
            "valor_remunerativo": ["valor_remunerativo"],
            "valor_no_remunerativo": ["valor_no_remunerativo"],
            "hs_trabajadas": ["hs_trabajadas"],
            "hs_con_justificacion_medica": ["hs_con_justificacion_medica", "no_trabajadas_c_justificativo_o_feriados"],
            "hs_feriado": ["hs_feriado"],
            "hs_extra": ["hs_extra"],
            "premios": ["premios"],
            "sueldo_bruto": ["sueldo_bruto", "bruto"],
            "sueldo_neto": ["sueldo_neto", "neto"],
        },
    },
    "comisiones": {
        "columns": ["id_comision", "fecha", "id_canal", "id_acreedor_etiqueta", "id_venta", "subtotal", "comision"],
        "aliases": {
            "id_comision": ["id_comision", "_rowNumber"],
            "fecha": ["fecha"],
            "id_canal": ["id_canal"],
            "id_acreedor_etiqueta": ["id_acreedor_etiqueta"],
            "id_venta": ["id_venta"],
            "subtotal": ["subtotal"],
            "comision": ["comision"],
        },
    },
    "inventarios": {
        "columns": ["id_inventario", "fecha", "turno", "id_empleado", "valor_total"],
        "aliases": {
            "id_inventario": ["id_inventario"],
            "fecha": ["fecha"],
            "turno": ["turno"],
            "id_empleado": ["id_empleado"],
            "valor_total": ["valor_total"],
        },
    },
    "detalle_inventarios": {
        "columns": ["id_detalle_inventario", "id_inventario", "id_item", "cantidad", "costo_unitario_usado", "valor_total"],
        "aliases": {
            "id_detalle_inventario": ["id_detalle_inventario", "_rowNumber"],
            "id_inventario": ["id_inventario"],
            "id_item": ["id_item"],
            "cantidad": ["cantidad", "cantidad_formulario"],
            "costo_unitario_usado": ["costo_unitario_usado"],
            "valor_total": ["valor_total"],
        },
    },
}


def main():
    parser = argparse.ArgumentParser(description="Build SunNutrition backend table cache from Excel workbooks.")
    parser.add_argument("--registry", default=str(REGISTRY_PATH))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    parser.add_argument("--workbook-dir", default="")
    args = parser.parse_args()

    registry = read_json(Path(args.registry))
    workbook_dir = Path(args.workbook_dir or registry.get("defaultWorkbookDir") or ".")
    output_path = Path(args.output)
    cache = {
        "version": registry.get("version", 1),
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "workbookDir": str(workbook_dir),
        "tables": {},
        "errors": [],
    }

    workbook_cache = {}
    for table in registry.get("tables", []):
        table_name = table["name"]
        workbook_name = table["workbook"]
        sheet_name = table["sheet"]
        workbook_path = workbook_dir / workbook_name

        if table.get("synthetic"):
            cache["tables"][table_name] = {
                "definition": table,
                "headers": [{"key": column, "label": column} for column in BACKEND_TABLE_PROJECTIONS.get(table_name, {}).get("columns", [])],
                "rows": [],
                "rowCount": 0,
                "source": {"synthetic": True},
            }
            continue

        try:
            workbook = workbook_cache.get(workbook_name)
            if workbook is None:
                if not workbook_path.exists():
                    raise FileNotFoundError(f"No se encontro {workbook_path}")
                workbook = load_workbook(workbook_path, read_only=True, data_only=True)
                workbook_cache[workbook_name] = workbook

            if sheet_name not in workbook.sheetnames:
                raise KeyError(f"No existe la solapa {sheet_name}")

            worksheet = workbook[sheet_name]
            extracted = extract_table(worksheet)
            source_rows = extracted["rows"]
            extracted = project_table(table_name, extracted)
            cache["tables"][table_name] = {
                "definition": table,
                "headers": extracted["headers"],
                "rows": extracted["rows"],
                "rowCount": len(extracted["rows"]),
                "_sourceRows": source_rows,
                "source": {
                    "workbook": workbook_name,
                    "sheet": sheet_name,
                    "headerRow": extracted["headerRow"],
                },
            }
        except Exception as error:
            cache["tables"][table_name] = {
                "definition": table,
                "headers": [],
                "rows": [],
                "rowCount": 0,
                "source": {"workbook": workbook_name, "sheet": sheet_name},
            }
            cache["errors"].append({
                "table": table_name,
                "workbook": workbook_name,
                "sheet": sheet_name,
                "error": str(error),
            })

    for workbook in workbook_cache.values():
        workbook.close()

    finalize_cache(cache)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    print(str(output_path))


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def extract_table(worksheet):
    header_row_index, headers = detect_header_row(worksheet)
    normalized_headers = dedupe_headers([normalize_header(header, index) for index, header in enumerate(headers)])
    rows = []

    for excel_row in worksheet.iter_rows(min_row=header_row_index + 1, values_only=True):
        values = list(excel_row[:len(headers)])
        if is_empty_row(values):
            continue

        record = {}
        raw_record = {}
        for index, header in enumerate(headers):
            raw_value = values[index] if index < len(values) else None
            clean_value = clean_cell(raw_value)
            record[normalized_headers[index]] = clean_value
            raw_record[str(header)] = clean_value

        if not meaningful_record(record):
            continue

        rows.append({
            **record,
            "_rowNumber": len(rows) + header_row_index + 1,
        })

    return {
        "headerRow": header_row_index,
        "headers": [
            {"key": normalized_headers[index], "label": str(headers[index])}
            for index in range(len(headers))
        ],
        "rows": rows,
    }


def project_table(table_name, extracted):
    projection = BACKEND_TABLE_PROJECTIONS.get(table_name)
    if not projection:
        return extracted

    projected_rows = []
    for row in extracted["rows"]:
        projected = {"_rowNumber": row.get("_rowNumber", "")}
        for column in projection["columns"]:
            projected[column] = value_from_aliases(row, projection["aliases"].get(column, [column]))
        projected_rows.append(finalize_projected_record(table_name, projected, row))

    return {
        **extracted,
        "headers": [{"key": column, "label": column} for column in projection["columns"]],
        "rows": projected_rows,
    }


def value_from_aliases(row, aliases):
    for alias in aliases:
        if alias in row:
            return row[alias]
    return ""


def finalize_projected_record(table_name, projected, source_row):
    if table_name == "acreedores":
        projected["_nombre_acreedor"] = source_row.get("nombre") or ""
        return projected

    if table_name == "recepciones":
        projected["_id_acreedor"] = source_row.get("id_acreedor") or ""
        projected["_etiqueta_gasto"] = source_row.get("tipo") or ""
        projected["_fecha_factura"] = source_row.get("fecha_factura") or ""
        projected["_tipo_factura"] = source_row.get("tipo_de_factura") or source_row.get("tipo_factura") or ""
        projected["_nro_factura"] = source_row.get("nro_factura") or ""
        projected["_subtotal"] = source_row.get("subtotal") or ""
        projected["_total"] = source_row.get("total") or ""
        return projected

    if table_name == "otros_gastos":
        projected["_id_acreedor"] = source_row.get("id_acreedor") or ""
        projected["_etiqueta_gasto"] = source_row.get("tipo") or ""
        projected["_fecha_factura"] = source_row.get("fecha_factura") or ""
        projected["_tipo_factura"] = source_row.get("tipo_de_factura") or source_row.get("tipo_factura") or ""
        projected["_nro_factura"] = source_row.get("nro_factura") or ""
        projected["_subtotal"] = source_row.get("subtotal") or ""
        projected["_total"] = source_row.get("total") or ""
        return projected

    if table_name == "entregas":
        projected["_id_acreedor"] = source_row.get("id_acreedor") or ""
        projected["_etiqueta_gasto"] = source_row.get("tipo") or "Logistica"
        projected["_fecha_factura"] = source_row.get("fecha_factura") or ""
        projected["_tipo_factura"] = source_row.get("tipo_de_factura") or source_row.get("tipo_factura") or ""
        projected["_nro_factura"] = source_row.get("nro_factura") or ""
        projected["_subtotal"] = source_row.get("subtotal") or ""
        projected["_total"] = source_row.get("total") or ""
        return projected

    if table_name == "egresos":
        projected["_id_acreedor"] = source_row.get("id_acreedor") or ""
        projected["_nombre_acreedor"] = source_row.get("nombre_acreedor") or ""
        projected["_etiqueta_gasto"] = source_row.get("tipo") or ""
        return projected

    if table_name == "detalle_pagos":
        projected["_nombre_acreedor"] = source_row.get("acreedor") or source_row.get("nombre_acreedor") or ""
        projected["_fecha_factura"] = source_row.get("fecha") or source_row.get("fecha_factura") or ""
        projected["_nro_factura"] = source_row.get("nro_factura") or ""
        return projected

    if table_name == "items":
        projected["_traduccion_ud_receta"] = source_row.get("traduccion_ud_receta") or ""
        return projected

    if table_name == "recetas":
        projected["_id_producto_resultado"] = source_row.get("id_producto") or ""
        projected["_tipo_componente"] = source_row.get("tipo") or ""
        return projected

    if table_name == "insumos":
        projected["_id_item"] = source_row.get("id_item") or ""
        return projected

    if table_name == "canales":
        projected["comision"] = projected.get("comision") or default_channel_commission(projected.get("id_canal"))
        return projected

    if table_name != "sueldos":
        return projected

    projected["fecha"] = payroll_month_to_date(projected.get("fecha"))
    projected["premios"] = sum_numbers(source_row.get("presentismo"), source_row.get("premio_pochocleras"), projected.get("premios"))
    projected["id_sueldo"] = projected.get("id_sueldo") or source_row.get("_rowNumber") or ""
    split_payroll_holiday_hours(projected)
    projected["_empleado_nombre"] = source_row.get("empleado") or source_row.get("nombre") or ""
    return projected


def assign_employee_contract_types(cache):
    for employee in cache.get("tables", {}).get("empleados", {}).get("rows", []):
        if employee.get("contratacion"):
            continue
        if not clean_text(employee.get("fecha_alta")) or clean_text(employee.get("fecha_baja")):
            continue
        normalized_name = normalize_lookup_text(employee.get("nombre_empleado") or employee.get("nombre"))
        if "alcaraz pablo" in normalized_name or "alcarez pablo" in normalized_name or "milciades ayala" in normalized_name:
            employee["contratacion"] = "Relacion de dependencia"
        elif "alexis kolln" in normalized_name:
            employee["contratacion"] = "Blue"
        else:
            employee["contratacion"] = "Cooperativa"


def finalize_cache(cache):
    expand_legacy_inventory_turn_tables(cache)
    assign_employee_contract_types(cache)
    resolve_creditor_origin_ids(cache)
    resolve_expense_creditor_tag_ids(cache)
    link_source_rows_to_expenses(cache)
    resolve_payment_expense_ids(cache)

    employee_by_name = {}
    for employee in cache.get("tables", {}).get("empleados", {}).get("rows", []):
        key = normalize_lookup_text(employee.get("nombre_empleado") or employee.get("nombre"))
        if key and employee.get("id_empleado") != "":
            employee_by_name[key] = employee.get("id_empleado")
    delivery_by_order = {}
    for detail in cache.get("tables", {}).get("entregas_detalle", {}).get("rows", []):
        order_id = str(detail.get("id_pedido") or "").strip()
        if order_id and order_id not in delivery_by_order:
            delivery_by_order[order_id] = detail.get("id_entrega") or ""
    item_meta_by_id = {}
    for item in cache.get("tables", {}).get("items", {}).get("rows", []):
        if item.get("id_item"):
            item_meta_by_id[str(item.get("id_item"))] = dict(item)
    item_origin_by_id = {}
    result_item_by_product_id = {}
    subproduct_item_ids = set()
    for product in cache.get("tables", {}).get("productos", {}).get("rows", []):
        item_id = str(product.get("id_item") or "").strip()
        if product.get("id_producto"):
            result_item_by_product_id[str(product.get("id_producto"))] = product.get("id_item") or ""
        if item_id:
            item_origin_by_id[item_id] = {"origen_tipo": "producto", "id_origen": product.get("id_producto") or ""}
    for recipe in cache.get("tables", {}).get("recetas", {}).get("rows", []):
        if normalize_origin_type(recipe.get("_tipo_componente")) == "subproducto" and recipe.get("id_item_componente"):
            subproduct_item_ids.add(str(recipe.get("id_item_componente")))
    for supply in cache.get("tables", {}).get("insumos", {}).get("rows", []):
        item_id = str(supply.get("_id_item") or "").strip()
        if item_id:
            item_origin_by_id[item_id] = {"origen_tipo": "insumo", "id_origen": supply.get("id_insumo") or ""}

    build_subproducts_table(cache, subproduct_item_ids, item_meta_by_id)

    for salary in cache.get("tables", {}).get("sueldos", {}).get("rows", []):
        if not salary.get("id_empleado"):
            salary["id_empleado"] = employee_by_name.get(normalize_lookup_text(salary.get("_empleado_nombre")), "")
        salary.pop("_empleado_nombre", None)

    for sale in cache.get("tables", {}).get("ventas", {}).get("rows", []):
        if not sale.get("id_entrega"):
            sale["id_entrega"] = delivery_by_order.get(str(sale.get("id_pedido") or "").strip(), "")

    for item in cache.get("tables", {}).get("items", {}).get("rows", []):
        item_id = str(item.get("id_item") or "").strip()
        origin = item_origin_by_id.get(item_id)
        if item_id in subproduct_item_ids:
            item["origen_tipo"] = "subproducto"
            item["id_origen"] = item.get("id_item")
        elif origin:
            item["origen_tipo"] = normalize_origin_type(origin.get("origen_tipo"))
            item["id_origen"] = origin.get("id_origen")
        else:
            item["origen_tipo"] = normalize_origin_type(item.get("origen_tipo"))

    for supply in cache.get("tables", {}).get("insumos", {}).get("rows", []):
        meta = item_meta_by_id.get(str(supply.get("_id_item") or "").strip())
        if meta and meta.get("_traduccion_ud_receta") != "":
            supply["cantidad_receta"] = meta.get("_traduccion_ud_receta")
        supply.pop("_id_item", None)

    dedupe_rows(cache, "insumos", "id_insumo")
    for recipe in cache.get("tables", {}).get("recetas", {}).get("rows", []):
        if not recipe.get("id_item_resultado"):
            recipe["id_item_resultado"] = result_item_by_product_id.get(str(recipe.get("_id_producto_resultado") or "").strip(), "")
        recipe.pop("_id_producto_resultado", None)
        recipe.pop("_tipo_componente", None)
    products = cache.get("tables", {}).get("productos")
    if products and "rows" in products:
        products["rows"] = [
            product for product in products.get("rows", [])
            if str(product.get("id_item") or "").strip() not in subproduct_item_ids
        ]
        products["rowCount"] = len(products["rows"])

    fill_empty_primary_ids(cache)
    value_backend_inventories(cache)
    cleanup_auxiliary_fields(cache)

    for table in cache.get("tables", {}).values():
        table.pop("_sourceRows", None)


def build_subproducts_table(cache, subproduct_item_ids, item_meta_by_id):
    cache.setdefault("tables", {})
    rows = []
    seen = set()

    def push_subproduct(row):
        subproduct_id = str(row.get("id_subproducto") or row.get("id_item") or "").strip()
        if not subproduct_id or subproduct_id in seen:
            return
        seen.add(subproduct_id)
        rows.append({
            "id_subproducto": subproduct_id,
            "id_item": row.get("id_item") or subproduct_id,
            "nombre_subproducto": row.get("nombre_subproducto") or f"Subproducto_{subproduct_id}",
            "ud_conteo": row.get("ud_conteo") or "",
        })

    for product in cache.get("tables", {}).get("productos", {}).get("rows", []):
        item_id = str(product.get("id_item") or "").strip()
        if item_id not in subproduct_item_ids:
            continue
        meta = item_meta_by_id.get(item_id, {})
        push_subproduct({
            "id_subproducto": item_id,
            "id_item": item_id,
            "nombre_subproducto": product.get("nombre_producto") or product.get("nombre") or "",
            "ud_conteo": meta.get("ud_conteo") or product.get("ud_conteo") or "",
        })

    for item_id in subproduct_item_ids:
        meta = item_meta_by_id.get(str(item_id), {})
        push_subproduct({
            "id_subproducto": item_id,
            "id_item": item_id,
            "nombre_subproducto": meta.get("nombre_subproducto") or "",
            "ud_conteo": meta.get("ud_conteo") or "",
        })

    cache["tables"]["subproductos"] = {
        "definition": {
            "name": "subproductos",
            "label": "Subproductos",
            "module": "inventario",
            "workbook": "Informacion.xlsx",
            "sheet": "Subproductos",
            "primaryKey": "id_subproducto",
            "synthetic": True,
        },
        "headers": [{"key": column, "label": column} for column in BACKEND_TABLE_PROJECTIONS["subproductos"]["columns"]],
        "rows": rows,
        "rowCount": len(rows),
        "source": {"syntheticFrom": ["productos", "items", "recetas"]},
    }


def resolve_creditor_origin_ids(cache):
    origin_maps = {
        "proveedor": rows_by_name(cache.get("tables", {}).get("proveedores", {}).get("rows", []), "nombre", "id_proveedor"),
        "empleado": rows_by_name(cache.get("tables", {}).get("empleados", {}).get("rows", []), "nombre_empleado", "id_empleado"),
        "canal": rows_by_name(cache.get("tables", {}).get("canales", {}).get("rows", []), "nombre", "id_canal"),
        "flete": rows_by_name(cache.get("tables", {}).get("fletes", {}).get("rows", []), "nombre_flete", "id_flete"),
    }

    for creditor in cache.get("tables", {}).get("acreedores", {}).get("rows", []):
        origin_type = creditor_origin_type_key(creditor.get("origen_tipo_acreedor"))
        creditor_name = normalize_lookup_text(creditor.get("_nombre_acreedor"))
        origin_id = origin_maps.get(origin_type, {}).get(creditor_name)
        if origin_id:
            creditor["origen_id_acreedor"] = origin_id
        creditor.pop("_nombre_acreedor", None)


def resolve_expense_creditor_tag_ids(cache):
    tags_by_name = {
        normalize_lookup_text(row.get("etiqueta")): clean_text(row.get("id_etiqueta"))
        for row in cache.get("tables", {}).get("etiquetas", {}).get("rows", [])
    }
    relations_by_creditor_tag = {}
    first_relation_by_creditor = {}
    relations = sorted(
        cache.get("tables", {}).get("acreedores_etiquetas", {}).get("rows", []),
        key=lambda row: numeric_sort_key(row.get("id_acreedor_etiqueta")),
    )

    for relation in relations:
        relation_id = clean_text(relation.get("id_acreedor_etiqueta"))
        creditor_id = clean_text(relation.get("id_acreedor"))
        tag_id = clean_text(relation.get("id_etiqueta"))
        if creditor_id and relation_id and creditor_id not in first_relation_by_creditor:
            first_relation_by_creditor[creditor_id] = relation_id
        if creditor_id and tag_id and relation_id:
            relations_by_creditor_tag[(creditor_id, tag_id)] = relation_id

    for table_name in ["recepciones", "otros_gastos", "entregas", "sueldos", "comisiones"]:
        for row in cache.get("tables", {}).get(table_name, {}).get("rows", []):
            if clean_text(row.get("id_acreedor_etiqueta")):
                continue
            creditor_id = clean_text(row.get("_id_acreedor")) or creditor_id_for_operational_row(cache, table_name, row)
            tag_id = tags_by_name.get(normalize_lookup_text(row.get("_etiqueta_gasto")), "")
            relation_id = relations_by_creditor_tag.get((creditor_id, tag_id)) or first_relation_by_creditor.get(creditor_id)
            if relation_id:
                row["id_acreedor_etiqueta"] = relation_id


def creditor_id_for_operational_row(cache, table_name, row):
    if table_name == "sueldos":
        return creditor_id_by_origin(cache, "empleado", row.get("id_empleado"))
    if table_name == "entregas":
        return creditor_id_by_origin(cache, "flete", row.get("id_flete"))
    if table_name == "comisiones":
        return creditor_id_by_origin(cache, "canal", row.get("id_canal"))
    if table_name == "otros_gastos":
        return clean_text(row.get("id_acreedor"))
    if table_name == "recepciones":
        purchase = row_by_id(cache.get("tables", {}).get("compras", {}).get("rows", []), "id_compra", row.get("id_compra"))
        return creditor_id_by_origin(cache, "proveedor", purchase.get("id_proveedor") if purchase else "")
    return ""


def creditor_id_by_origin(cache, origin_type, origin_id):
    expected_type = creditor_origin_type_key(origin_type)
    expected_id = clean_text(origin_id)
    if not expected_type or not expected_id:
        return ""
    for creditor in cache.get("tables", {}).get("acreedores", {}).get("rows", []):
        if creditor_origin_type_key(creditor.get("origen_tipo_acreedor")) != expected_type:
            continue
        if clean_text(creditor.get("origen_id_acreedor")) == expected_id:
            return clean_text(creditor.get("id_acreedor"))
    return ""


def row_by_id(rows, id_column, value):
    expected = clean_text(value)
    if not expected:
        return None
    for row in rows or []:
        if clean_text(row.get(id_column)) == expected:
            return row
    return None


def numeric_sort_key(value):
    try:
        return int(float(str(value).replace(",", ".")))
    except (TypeError, ValueError):
        return 999999999


def resolve_payment_expense_ids(cache):
    expense_by_key = {}
    for expense in cache.get("tables", {}).get("egresos", {}).get("rows", []):
        key = payment_expense_key(expense.get("_nombre_acreedor"), expense.get("fecha_factura"), expense.get("nro_factura"))
        if key and key not in expense_by_key:
            expense_by_key[key] = expense.get("id_egreso")

    for detail in cache.get("tables", {}).get("detalle_pagos", {}).get("rows", []):
        if clean_text(detail.get("id_egreso")):
            continue
        key = payment_expense_key(detail.get("_nombre_acreedor"), detail.get("_fecha_factura"), detail.get("_nro_factura"))
        expense_id = expense_by_key.get(key)
        if expense_id:
            detail["id_egreso"] = expense_id


def payment_expense_key(creditor_name, invoice_date, invoice_number):
    values = [
        normalize_lookup_text(creditor_name),
        normalize_date_key(invoice_date),
        normalize_invoice_number(invoice_number),
    ]
    if any(not value for value in values):
        return ""
    return "|".join(values)


def link_source_rows_to_expenses(cache):
    tables = cache.get("tables", {})
    expense_by_legacy_origin = {}
    expense_by_operational_key = {}

    for expense in tables.get("egresos", {}).get("rows", []):
        expense_id = clean_text(expense.get("id_egreso"))
        legacy_type = expense_source_type_key(expense.get("origen_tipo"))
        legacy_id = clean_text(expense.get("origen_id"))
        if expense_id and legacy_type and legacy_id:
            expense_by_legacy_origin[f"{legacy_type}:{legacy_id}"] = expense_id

        match_key = expense_operational_match_key(expense)
        if expense_id and match_key and match_key not in expense_by_operational_key:
            expense_by_operational_key[match_key] = expense_id

        expense.pop("origen_tipo", None)
        expense.pop("origen_id", None)

    for config in expense_source_configs():
        for row in tables.get(config["table"], {}).get("rows", []):
            if clean_text(row.get("id_egreso")):
                continue
            source_id = clean_text(row.get(config["id_column"]))
            legacy_match = expense_by_legacy_origin.get(f"{expense_source_type_key(config['type'])}:{source_id}")
            if legacy_match:
                row["id_egreso"] = legacy_match
                continue
            operational_match = expense_by_operational_key.get(expense_operational_match_key(row))
            if operational_match:
                row["id_egreso"] = operational_match


def expense_source_configs():
    return [
        {"table": "recepciones", "type": "recepciones", "id_column": "id_recepcion"},
        {"table": "otros_gastos", "type": "otros_gastos", "id_column": "id_otros_gastos"},
        {"table": "sueldos", "type": "sueldos", "id_column": "id_sueldo"},
        {"table": "entregas", "type": "logistica", "id_column": "id_entrega"},
        {"table": "comisiones", "type": "comisiones", "id_column": "id_comision"},
    ]


def expense_source_type_key(value):
    text = normalize_lookup_text(value)
    if text in ("entregas", "logistica"):
        return "logistica"
    if text in ("otros gastos", "otros_gastos", "otrosgastos"):
        return "otros_gastos"
    return text


def expense_operational_match_key(row):
    values = [
        clean_text(row.get("_id_acreedor")),
        normalize_date_key(row.get("_fecha_factura") or row.get("fecha_factura")),
        normalize_lookup_text(row.get("_tipo_factura") or row.get("tipo_factura")),
        normalize_invoice_number(row.get("_nro_factura") or row.get("nro_factura")),
        normalize_amount_key(row.get("_total") or row.get("total")),
    ]
    if any(not value for value in values):
        return ""
    return "|".join(values)


def clean_text(value):
    return str(value or "").strip()


def normalize_date_key(value):
    text = clean_text(value)
    if not text:
        return ""
    iso_match = re.match(r"^(\d{4})-(\d{2})-(\d{2})", text)
    if iso_match:
        return f"{iso_match.group(1)}-{iso_match.group(2)}-{iso_match.group(3)}"
    local_match = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})", text)
    if local_match:
        year = local_match.group(3)
        if len(year) == 2:
            year = f"20{year}"
        return f"{year}-{local_match.group(2).zfill(2)}-{local_match.group(1).zfill(2)}"
    return normalize_lookup_text(text)


def normalize_invoice_number(value):
    text = clean_text(value)
    if not text:
        return ""
    try:
        return str(int(float(text.replace(",", "."))))
    except ValueError:
        return normalize_lookup_text(text)


def normalize_amount_key(value):
    text = clean_text(value)
    if not text:
        return ""
    normalized = re.sub(r"\.(?=\d{3}(\D|$))", "", text.replace("$", "").replace(" ", "")).replace(",", ".")
    try:
        number = float(normalized)
    except ValueError:
        return normalize_lookup_text(text)
    return str(round(number, 2))


def cleanup_auxiliary_fields(cache):
    auxiliary_columns = [
        "_nombre_acreedor",
        "_traduccion_ud_receta",
    ]
    for table in cache.get("tables", {}).values():
        for row in table.get("rows", []):
            for column in auxiliary_columns:
                row.pop(column, None)


def rows_by_name(rows, name_column, id_column):
    result = {}
    for row in rows:
        key = normalize_lookup_text(row.get(name_column))
        row_id = str(row.get(id_column) or "").strip()
        if key and row_id and key not in result:
            result[key] = row_id
    return result


def creditor_origin_type_key(value):
    text = normalize_lookup_text(value)
    if "proveedor" in text:
        return "proveedor"
    if "empleado" in text:
        return "empleado"
    if "canal" in text:
        return "canal"
    if "flete" in text:
        return "flete"
    return ""


def fill_empty_primary_ids(cache):
    for table_name, table in cache.get("tables", {}).items():
        projection = BACKEND_TABLE_PROJECTIONS.get(table_name, {})
        columns = projection.get("columns") or []
        primary_column = columns[0] if columns else ""
        rows = table.get("rows") or []
        if not primary_column.startswith("id_") or not rows:
            continue

        primary_id_is_empty = all(str(row.get(primary_column) or "").strip() == "" for row in rows)
        if not primary_id_is_empty:
            continue

        for index, row in enumerate(rows, start=1):
            row[primary_column] = index


def value_backend_inventories(cache):
    inventory_table = cache.get("tables", {}).get("inventarios")
    detail_table = cache.get("tables", {}).get("detalle_inventarios")
    if not inventory_table or not detail_table:
        return

    inventories_by_id = rows_by_id(inventory_table.get("rows", []), "id_inventario")
    details_by_inventory = group_rows_by_id(detail_table.get("rows", []), "id_inventario")
    item_costs_by_date = {}

    for inventory in inventory_table.get("rows", []):
        inventory["valor_total"] = 0

    sorted_inventories = sorted(
        inventory_table.get("rows", []),
        key=lambda inventory: (normalize_date_key(inventory.get("fecha")), turn_rank(inventory.get("turno"))),
    )

    for inventory in sorted_inventories:
        inventory_id = str(inventory.get("id_inventario") or "").strip()
        inventory_date = normalize_date_key(inventory.get("fecha"))
        if inventory_date not in item_costs_by_date:
            item_costs_by_date[inventory_date] = inventory_item_cost_map(cache, inventory_date)
        item_costs = item_costs_by_date[inventory_date]
        inventory_value = 0

        for detail in details_by_inventory.get(inventory_id, []):
            detail["costo_unitario_usado"] = 0
            detail["valor_total"] = 0

            unit_cost = item_costs.get(str(detail.get("id_item") or "").strip())
            if not unit_cost:
                continue

            quantity = parse_number(detail.get("cantidad"))
            line_value = quantity * unit_cost
            detail["costo_unitario_usado"] = round_money(unit_cost)
            detail["valor_total"] = round_money(line_value)
            inventory_value += line_value

        stored_inventory = inventories_by_id.get(inventory_id, inventory)
        stored_inventory["valor_total"] = round_money(inventory_value)

    inventory_table["headers"] = [{"key": column, "label": column} for column in BACKEND_TABLE_PROJECTIONS["inventarios"]["columns"]]
    detail_table["headers"] = [{"key": column, "label": column} for column in BACKEND_TABLE_PROJECTIONS["detalle_inventarios"]["columns"]]


def inventory_item_cost_map(cache, date_iso):
    items_by_id = rows_by_id(cache.get("tables", {}).get("items", {}).get("rows", []), "id_item")
    supplies_by_id = rows_by_id(cache.get("tables", {}).get("insumos", {}).get("rows", []), "id_insumo")
    recipes_by_result = group_rows_by_id(cache.get("tables", {}).get("recetas", {}).get("rows", []), "id_item_resultado")
    supply_costs = latest_supply_cost_map(cache, date_iso)
    recipe_unit_memo = {}
    count_unit_costs = {}
    visiting = set()

    def count_to_recipe(item):
        item_translation = parse_number(item.get("_traduccion_ud_receta") if item else "")
        if item_translation > 0:
            return item_translation

        if normalize_lookup_text(item.get("origen_tipo") if item else "") == "insumo":
            supply = supplies_by_id.get(str(item.get("id_origen") or "").strip(), {})
            supply_translation = parse_number(supply.get("cantidad_receta"))
            if supply_translation > 0:
                return supply_translation

        return 1

    def cost_per_recipe_unit(item_id):
        normalized_item_id = str(item_id or "").strip()
        if not normalized_item_id:
            return 0
        if normalized_item_id in recipe_unit_memo:
            return recipe_unit_memo[normalized_item_id]
        if normalized_item_id in visiting:
            return 0

        visiting.add(normalized_item_id)
        item = items_by_id.get(normalized_item_id, {})
        cost = 0

        if normalize_lookup_text(item.get("origen_tipo")) == "insumo":
            supply_id = str(item.get("id_origen") or "").strip()
            supply_cost = supply_costs.get(supply_id)
            if supply_cost and supply_cost.get("unit_cost"):
                cost = supply_cost["unit_cost"]
        else:
            for recipe in recipes_by_result.get(normalized_item_id, []):
                component_quantity = parse_number(recipe.get("cantidad_componente"))
                component_cost = cost_per_recipe_unit(recipe.get("id_item_componente"))
                cost += component_quantity * component_cost

        visiting.remove(normalized_item_id)
        recipe_unit_memo[normalized_item_id] = cost
        return cost

    for item_id, item in items_by_id.items():
        cost = cost_per_recipe_unit(item_id) * count_to_recipe(item)
        if cost > 0:
            count_unit_costs[item_id] = cost

    return count_unit_costs


def latest_supply_cost_map(cache, date_iso):
    purchases_by_id = rows_by_id(cache.get("tables", {}).get("compras", {}).get("rows", []), "id_compra")
    receptions_by_purchase = group_rows_by_id(cache.get("tables", {}).get("recepciones", {}).get("rows", []), "id_compra")
    supplier_items_by_id = rows_by_id(cache.get("tables", {}).get("insumos_proveedores", {}).get("rows", []), "id_insumos_proveedores")
    candidates_by_supply = {}

    for detail in cache.get("tables", {}).get("detalle_compras", {}).get("rows", []):
        purchase = purchases_by_id.get(str(detail.get("id_compra") or "").strip())
        if not purchase:
            continue

        reception_dates = sorted(
            filter(None, (normalize_date_key(row.get("fecha_recepcion")) for row in receptions_by_purchase.get(str(detail.get("id_compra") or "").strip(), [])))
        )
        purchase_date = (
            reception_dates[-1] if reception_dates else
            normalize_date_key(purchase.get("fecha_pedido")) or
            normalize_date_key(purchase.get("fecha_entrega_prevista"))
        )
        if not purchase_date:
            continue

        supplier_item = supplier_items_by_id.get(str(detail.get("id_insumos_proveedores") or "").strip())
        if not supplier_item:
            continue

        supply_id = str(supplier_item.get("id_insumo") or "").strip()
        supplier_unit_quantity = parse_number(supplier_item.get("cantidad_proveedor")) or 1
        supplier_unit_price = parse_number(supplier_item.get("precio"))
        if not supply_id or not supplier_unit_price:
            continue

        unit_cost = supplier_unit_price / supplier_unit_quantity
        candidates_by_supply.setdefault(supply_id, []).append({
            "date": purchase_date,
            "unit_cost": unit_cost,
        })

    selected_by_supply = {}
    for supply_id, candidates in candidates_by_supply.items():
        sorted_candidates = sorted(candidates, key=lambda candidate: candidate["date"])
        previous_candidates = [candidate for candidate in sorted_candidates if not date_iso or candidate["date"] <= date_iso]
        selected_by_supply[supply_id] = previous_candidates[-1] if previous_candidates else sorted_candidates[0]

    return selected_by_supply


def rows_by_id(rows, id_column):
    result = {}
    for row in rows or []:
        row_id = str(row.get(id_column) or "").strip()
        if row_id and row_id not in result:
            result[row_id] = row
    return result


def value_is_blank(value):
    return value is None or str(value).strip() == ""


def group_rows_by_id(rows, id_column):
    result = {}
    for row in rows or []:
        row_id = str(row.get(id_column) or "").strip()
        if not row_id:
            continue
        result.setdefault(row_id, []).append(row)
    return result


def turn_rank(value):
    text = normalize_lookup_text(value)
    if "madrugada" in text:
        return 1
    if "manana" in text or "mañana" in text:
        return 2
    if "tarde" in text:
        return 3
    return 0


def round_money(value):
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return 0


def expand_legacy_inventory_turn_tables(cache):
    inventory_table = cache.get("tables", {}).get("inventarios")
    detail_table = cache.get("tables", {}).get("detalle_inventarios")
    source_inventories = inventory_table.get("_sourceRows", []) if inventory_table else []
    source_details = detail_table.get("_sourceRows", []) if detail_table else []

    if not inventory_table or not detail_table:
        return
    if not is_legacy_inventory_turn_source(source_inventories, source_details):
        return

    details_by_inventory = {}
    for detail in source_details:
        inventory_id = str(value_from_aliases(detail, ["id_inventario"]) or "").strip()
        if not inventory_id:
            continue
        details_by_inventory.setdefault(inventory_id, []).append(detail)

    migrated_inventories = []
    migrated_details = []
    next_inventory_id = 1
    next_detail_id = 1

    sorted_inventories = sorted(
        [inventory for inventory in source_inventories if str(value_from_aliases(inventory, ["id_inventario"]) or "").strip()],
        key=lambda inventory: numeric_sort_key(value_from_aliases(inventory, ["id_inventario"])),
    )

    for inventory in sorted_inventories:
        legacy_inventory_id = str(value_from_aliases(inventory, ["id_inventario"]) or "").strip()
        legacy_details = details_by_inventory.get(legacy_inventory_id, [])

        turn_groups = []
        for turn in legacy_inventory_turn_definitions():
            details_for_turn = []
            for detail in legacy_details:
                item_id = value_from_aliases(detail, ["id_item"])
                quantity = clean_legacy_inventory_quantity(value_from_aliases(detail, turn["aliases"]))
                if str(item_id or "").strip() and quantity != "":
                    details_for_turn.append({"id_item": item_id, "cantidad": quantity})
            if details_for_turn:
                turn_groups.append({"turn": turn, "details": details_for_turn})

        single_turn_day = len(turn_groups) == 1

        for group in turn_groups:
            inventory_id = next_inventory_id
            migrated_inventories.append({
                "id_inventario": inventory_id,
                "fecha": value_from_aliases(inventory, ["fecha"]),
                "turno": "Mañana" if single_turn_day else group["turn"]["label"],
                "id_empleado": value_from_aliases(inventory, ["id_empleado", "empleado"]),
                "valor_total": 0,
            })

            for detail in group["details"]:
                migrated_details.append({
                    "id_detalle_inventario": next_detail_id,
                    "id_inventario": inventory_id,
                    "id_item": detail["id_item"],
                    "cantidad": detail["cantidad"],
                    "costo_unitario_usado": 0,
                    "valor_total": 0,
                })
                next_detail_id += 1

            next_inventory_id += 1

    inventory_table["headers"] = [{"key": column, "label": column} for column in BACKEND_TABLE_PROJECTIONS["inventarios"]["columns"]]
    inventory_table["rows"] = migrated_inventories
    inventory_table["rowCount"] = len(migrated_inventories)
    inventory_table.setdefault("source", {})["transformedFromLegacyTurns"] = True

    detail_table["headers"] = [{"key": column, "label": column} for column in BACKEND_TABLE_PROJECTIONS["detalle_inventarios"]["columns"]]
    detail_table["rows"] = migrated_details
    detail_table["rowCount"] = len(migrated_details)
    detail_table.setdefault("source", {})["transformedFromLegacyTurns"] = True


def is_legacy_inventory_turn_source(source_inventories, source_details):
    if not source_inventories or not source_details:
        return False
    has_new_turn = any(str(inventory.get("turno") or "").strip() for inventory in source_inventories)
    has_legacy_columns = any(
        alias in detail
        for detail in source_details
        for turn in legacy_inventory_turn_definitions()
        for alias in turn["aliases"]
    )
    return not has_new_turn and has_legacy_columns


def legacy_inventory_turn_definitions():
    return [
        {"label": "Madrugada", "aliases": ["turno_madrugadacf", "turno_madrugada_cf", "turno_madrugada"]},
        {"label": "Mañana", "aliases": ["turno_mananacf", "turno_manana_cf", "turno_manana", "turno_mañanacf"]},
        {"label": "Tarde", "aliases": ["cantidad_formulario", "turno_tardecf", "turno_tarde_cf", "turno_tarde"]},
    ]


def clean_legacy_inventory_quantity(value):
    text = str(value or "").strip()
    if not text or re.match(r"^[-–—]+$", text):
        return ""
    return text.replace(",", ".")


def numeric_sort_key(value):
    try:
        return float(str(value or "").replace(",", "."))
    except ValueError:
        return float("inf")


def payroll_month_to_date(value):
    text = str(value or "").strip()
    match = re.match(r"^(\d{4})(\d{2})(?:\.0)?$", text)
    if not match:
        return value or ""
    return f"{match.group(1)}-{match.group(2)}-01"


def split_payroll_holiday_hours(payroll):
    if clean_text(payroll.get("hs_feriado")):
        return
    combined_hours = parse_number(payroll.get("hs_con_justificacion_medica"))
    if combined_hours is None or combined_hours <= 0:
        payroll["hs_feriado"] = ""
        return

    holiday_hours = min(combined_hours, argentina_business_holiday_count_in_month(payroll.get("fecha")) * 8)
    payroll["hs_feriado"] = format_number(holiday_hours)
    payroll["hs_con_justificacion_medica"] = format_number(combined_hours - holiday_hours)


def argentina_business_holiday_count_in_month(month_date):
    match = re.match(r"^(\d{4})-(\d{2})-", str(month_date or ""))
    if not match:
        return 0
    year = int(match.group(1))
    month = int(match.group(2))
    return sum(
        1
        for holiday in argentina_national_holidays(year)
        if holiday.month == month and holiday.weekday() < 5
    )


def argentina_national_holidays(year):
    easter = easter_sunday(year)
    holidays = [
        date(year, 1, 1),
        easter + date_delta(-48),
        easter + date_delta(-47),
        date(year, 3, 24),
        date(year, 4, 2),
        easter + date_delta(-2),
        date(year, 5, 1),
        date(year, 5, 25),
        transfer_argentina_holiday(date(year, 6, 17)),
        date(year, 6, 20),
        date(year, 7, 9),
        transfer_argentina_holiday(date(year, 8, 17)),
        transfer_argentina_holiday(date(year, 10, 12)),
        transfer_argentina_holiday(date(year, 11, 20)),
        date(year, 12, 8),
        date(year, 12, 25),
    ]
    return list(dict.fromkeys(holidays))


def transfer_argentina_holiday(day):
    weekday = day.weekday()
    if weekday in (1, 2):
        return day + date_delta(-weekday)
    if weekday in (3, 4):
        return day + date_delta(7 - weekday)
    return day


def easter_sunday(year):
    a = year % 19
    b = year // 100
    c = year % 100
    d = b // 4
    e = b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i = c // 4
    k = c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    month = (h + l - 7 * m + 114) // 31
    day = ((h + l - 7 * m + 114) % 31) + 1
    return date(year, month, day)


def date_delta(days):
    from datetime import timedelta
    return timedelta(days=days)


def parse_number(value):
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[^\d,.-]", "", str(value or "").strip())
    if not text:
        return 0
    has_comma = "," in text
    has_dot = "." in text
    if has_comma and has_dot:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif has_comma:
        text = text.replace(".", "").replace(",", ".")
    else:
        parts = text.split(".")
        if len(parts) > 2:
            text = "".join(parts[:-1]) + "." + parts[-1]
    try:
        return float(text)
    except ValueError:
        return 0


def format_number(value):
    if value is None:
        return ""
    rounded = round(value, 2)
    return int(rounded) if float(rounded).is_integer() else rounded


def sum_numbers(*values):
    total = 0
    for value in values:
        try:
            total += float(str(value or "").replace(",", "."))
        except ValueError:
            continue
    return int(total) if total and total.is_integer() else (total or "")


def default_channel_commission(channel_id):
    commissions = {
        1: 0.05,
        2: 0.06,
        3: 0.05,
        4: 0,
        5: 0.06,
        6: 0.05,
        7: 0.05,
    }
    try:
        key = int(float(str(channel_id or "").replace(",", ".")))
    except ValueError:
        return ""
    return commissions.get(key, "")


def normalize_lookup_text(value):
    text = str(value or "").strip().lower()
    accents = str.maketrans("áéíóúñü", "aeiounu")
    return re.sub(r"\s+", "", text.translate(accents))


def normalize_origin_type(value):
    text = str(value or "").strip().lower()
    if "subproducto" in text:
        return "subproducto"
    if "producto" in text:
        return "producto"
    if "insumo" in text:
        return "insumo"
    return text


def dedupe_rows(cache, table_name, key):
    table = cache.get("tables", {}).get(table_name)
    if not table or "rows" not in table:
        return
    seen = set()
    rows = []
    for row in table.get("rows", []):
        value = str(row.get(key) or "").strip()
        if value:
            if value in seen:
                continue
            seen.add(value)
        rows.append(row)
    table["rows"] = rows
    table["rowCount"] = len(rows)


def detect_header_row(worksheet, scan_limit=20):
    best = None
    for row_index, row in enumerate(worksheet.iter_rows(min_row=1, max_row=scan_limit, values_only=True), start=1):
        values = [clean_header(cell) for cell in row]
        non_empty = [value for value in values if value]
        if len(non_empty) < 2:
            continue
        score = header_score(non_empty)
        if best is None or score > best["score"]:
            best = {"row": row_index, "values": values, "score": score}

    if best is None:
        return 1, []

    headers = trim_trailing_empty(best["values"])
    return best["row"], headers


def header_score(values):
    score = len(values)
    for value in values:
        text = str(value)
        if re.search(r"\b(id|fecha|nombre|tipo|monto|total|cantidad|cliente|proveedor|acreedor)\b", text, re.I):
            score += 5
        if text.startswith("=") or "COMPUTED_VALUE" in text:
            score -= 6
    return score


def normalize_header(value, index):
    text = clean_header(value)
    if not text:
        return f"columna_{index + 1}"
    text = text.strip().lower()
    replacements = str.maketrans("áéíóúñü", "aeiounu")
    text = text.translate(replacements)
    text = re.sub(r"[^a-z0-9]+", "_", text)
    text = re.sub(r"^_+|_+$", "", text)
    return text or f"columna_{index + 1}"


def dedupe_headers(headers):
    seen = {}
    result = []
    for header in headers:
        count = seen.get(header, 0)
        seen[header] = count + 1
        result.append(header if count == 0 else f"{header}_{count + 1}")
    return result


def clean_header(value):
    if value is None:
        return ""
    if isinstance(value, ArrayFormula):
        return str(value)
    return str(value).strip()


def clean_cell(value):
    if value is None:
        return ""
    if isinstance(value, ArrayFormula):
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    if isinstance(value, str):
        return value.strip()
    return value


def trim_trailing_empty(values):
    values = list(values)
    while values and not values[-1]:
        values.pop()
    return values


def is_empty_row(values):
    return all(clean_cell(value) in ("", None) for value in values)


def meaningful_record(record):
    values = [value for key, value in record.items() if not key.startswith("columna_")]
    return any(value not in ("", None) for value in values)


if __name__ == "__main__":
    main()
